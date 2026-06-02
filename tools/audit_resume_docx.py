from pathlib import Path
import sys
from zipfile import ZipFile

from docx import Document


sys.stdout.reconfigure(encoding="utf-8", errors="replace")
DOCX = Path(r"C:\Users\Lenovo\OneDrive\Desktop\何文旭-河海大学-后端开发-15737679913-模板优化版.docx")


def main():
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    doc = Document(DOCX)
    sections = doc.sections
    text = "\n".join(p.text for p in doc.paragraphs)
    required = ["何文旭", "教育经历", "项目经验", "Enterprise HR AI Agent", "pet-dispatch-pro", "专业技能"]
    missing = [item for item in required if item not in text]
    if missing:
        raise AssertionError(f"missing required text: {missing}")

    section = sections[0]
    print(f"FILE={DOCX}")
    print(f"SIZE={DOCX.stat().st_size}")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)}")
    print(
        "PAGE_CM="
        f"{section.page_width.cm:.2f}x{section.page_height.cm:.2f} "
        f"MARGINS_IN={section.top_margin.inches:.2f},"
        f"{section.right_margin.inches:.2f},"
        f"{section.bottom_margin.inches:.2f},"
        f"{section.left_margin.inches:.2f}"
    )

    with ZipFile(DOCX) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
        has_heading_rule = 'w:bottom w:val="single"' in document_xml or "<w:bottom" in document_xml
        has_tag_shading = 'w:fill="DFF2FF"' in document_xml
        print(f"HEADING_RULE={has_heading_rule} TAG_SHADING={has_tag_shading}")

    print("AUDIT=OK")


if __name__ == "__main__":
    main()
