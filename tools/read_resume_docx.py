from pathlib import Path
import sys

from docx import Document


INPUT = Path(r"C:\Users\Lenovo\OneDrive\Desktop\何文旭-河海大学-后端开发-15737679913.docx")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def iter_table_rows(table):
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append(text)
        yield " | ".join(cells)


def main():
    doc = Document(INPUT)
    print(f"FILE: {INPUT}")
    print("\n# PARAGRAPHS")
    for i, p in enumerate(doc.paragraphs, 1):
        text = p.text.strip()
        if text:
            print(f"{i:03d}: {text}")
    print("\n# TABLES")
    for ti, table in enumerate(doc.tables, 1):
        print(f"TABLE {ti}")
        for row in iter_table_rows(table):
            if row.strip(" |"):
                print(row)


if __name__ == "__main__":
    main()
