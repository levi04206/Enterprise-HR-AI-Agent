from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


DESKTOP = Path(r"C:\Users\Lenovo\OneDrive\Desktop")
OUTPUT = DESKTOP / "何文旭-河海大学-后端开发-15737679913-模板优化版.docx"

FONT = "Microsoft YaHei"
INK = RGBColor(0x27, 0x31, 0x3B)
DARK = RGBColor(0x13, 0x23, 0x34)
MUTED = RGBColor(0x34, 0x49, 0x5E)
TAG_BLUE = RGBColor(0x16, 0x9B, 0xD5)
RULE = "262626"


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_font(paragraph, size=9.75, color=INK):
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color)


def set_paragraph_spacing(paragraph, before=0, after=1.8, line=1.12):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line


def add_bottom_border(paragraph, color=RULE, size="12"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_tag_run(paragraph, text):
    run = paragraph.add_run(text)
    set_run_font(run, size=8.8, bold=True, color=TAG_BLUE)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "DFF2FF")
    run._r.get_or_add_rPr().append(shading)
    return run


def add_section(doc, title):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=5.5, after=2.2, line=1.0)
    run = p.add_run(title)
    set_run_font(run, size=11.3, bold=True, color=RGBColor(0x14, 0x1A, 0x22))
    add_bottom_border(p)
    return p


def add_entry_header(doc, left_parts: Iterable[tuple[str, bool]], date: str = ""):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=1.4, after=0.6, line=1.0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.27), WD_ALIGN_PARAGRAPH.RIGHT)
    for text, bold in left_parts:
        run = p.add_run(text)
        set_run_font(run, size=9.8, bold=bold, color=RGBColor(0x1F, 0x29, 0x33))
    if date:
        run = p.add_run("\t" + date)
        set_run_font(run, size=9.3, bold=True, color=MUTED)
    return p


def add_label_line(doc, label, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=1.2, line=1.08)
    r = p.add_run(label)
    set_run_font(r, size=9.3, bold=True, color=DARK)
    r = p.add_run(text)
    set_run_font(r, size=9.3, color=INK)
    return p


def add_bullets(doc, bullets):
    for item in bullets:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, before=0, after=1.4, line=1.08)
        p.paragraph_format.left_indent = Inches(0.20)
        p.paragraph_format.first_line_indent = Inches(-0.10)
        run = p.add_run(item)
        set_run_font(run, size=9.15, color=INK)


def add_skill_line(doc, label, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=1.0, line=1.05)
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.10)
    r = p.add_run(label)
    set_run_font(r, size=9.05, bold=True, color=DARK)
    r = p.add_run(text)
    set_run_font(r, size=9.05, color=INK)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.38)
    section.bottom_margin = Inches(0.38)
    section.left_margin = Inches(0.50)
    section.right_margin = Inches(0.50)
    section.header_distance = Inches(0.20)
    section.footer_distance = Inches(0.20)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.75)
    normal.font.color.rgb = INK

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(9.15)


def build():
    doc = Document()
    configure_document(doc)

    # Header: Markdown template maps H1 to centered name and first paragraph to contact info.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=0.5, line=1.0)
    r = p.add_run("何文旭")
    set_run_font(r, size=15.8, bold=True, color=DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=0, after=5.0, line=1.0)
    r = p.add_run("Java后端开发实习生 | 15737679913 | 15737679913@163.com | 河海大学（211）硕士在读")
    set_run_font(r, size=9.45, color=INK)

    add_section(doc, "教育经历")
    add_entry_header(
        doc,
        [("河海大学（211） - 硕士", True), ("    土木与交通学院 / 岩土工程", False)],
        "2025.09 - 2028.06",
    )
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=1.0, line=1.0)
    add_tag_run(p, " 211高校 ")
    r = p.add_run("  具备工程背景与后端开发项目经验，目标岗位：Java 后端开发实习。")
    set_run_font(r, size=9.0, color=INK)

    add_entry_header(
        doc,
        [("河南理工大学 - 本科", True), ("    土木工程学院 / 工程管理", False)],
        "2021.09 - 2025.06",
    )

    add_section(doc, "项目经验")
    add_entry_header(
        doc,
        [("一、Enterprise HR AI Agent 企业 HR 智能助理系统", True), ("（个人项目）", False)],
        "2026.05 - 2026.06",
    )
    add_label_line(
        doc,
        "项目描述：",
        "面向企业 HR 场景的智能助理系统，覆盖员工管理、请假审批、知识库问答、会话持久化、工具调用和观测日志，支持管理员通过自然语言查询员工数量与明细。",
    )
    add_label_line(
        doc,
        "技术栈：",
        "Java 21、Spring Boot、Spring AI、MyBatis-Plus、MySQL、Flyway、SSE、RAG、Function Calling、JUnit 5、HTML/CSS/JavaScript。",
    )
    add_bullets(
        doc,
        [
            "实现员工 CRUD、关键词/ID 查询、请假提交、管理员审批/驳回及年假余额扣减，并通过全局异常处理向前端返回明确错误信息。",
            "基于 Spring AI ChatClient 实现 HR 智能对话，支持 SSE 流式响应、多轮会话与消息持久化，提升对话可追踪性。",
            "构建 RAG 知识库流程，支持 PDF/TXT 上传、文档解析、文本切分、向量化检索与命中片段记录，用于企业制度问答。",
            "接入 Function Calling，封装员工数量、员工详情、年假余额等工具；补充工具调用日志、RAG 命中日志和核心接口测试，便于调试和演示。",
        ],
    )

    add_entry_header(
        doc,
        [("二、pet-dispatch-pro 宠物上门服务调度系统", True), ("（个人独立开发）", False)],
        "2026.03 - 2026.05",
    )
    add_label_line(
        doc,
        "项目描述：",
        "面向宠物上门服务场景的调度平台，覆盖验证码登录、宠物档案、订单创建、宠托师接单、履约打卡、评价与管理员审核的完整闭环。",
    )
    add_label_line(
        doc,
        "技术栈：",
        "Java 21、Spring Boot、MyBatis-Plus、MySQL、Redis、JWT、Spring Validation、Spring AOP、阿里云 OSS、JUnit 5、MockMvc。",
    )
    add_bullets(
        doc,
        [
            "使用 DTO/VO、Spring Validation 与全局异常处理规范接口出入参，基于 Redis + JWT + 拦截器实现用户、宠托师、管理员三端权限隔离。",
            "围绕订单调度链路设计状态机，支持创建、支付、抢单、履约、取消和评价流程；通过事务、条件更新和乐观锁解决并发抢单冲突。",
            "基于 Spring AOP 实现接口幂等防重复提交与核心操作异步审计，降低重复请求和关键链路不可追踪风险。",
            "使用 Redis ZSET 重构高评分宠托师推荐榜单，并通过 JUnit/MockMvc 覆盖核心业务分支与权限边界。",
        ],
    )

    add_section(doc, "专业技能")
    add_skill_line(doc, "Java：", "熟悉集合、IO、异常、面向对象、常用设计模式；理解 JVM 内存模型、类加载机制、GC 与 JUC 并发基础。")
    add_skill_line(doc, "框架：", "熟悉 Spring Boot、Spring MVC、MyBatis/MyBatis-Plus，理解 IOC、AOP、Bean 生命周期与分层开发。")
    add_skill_line(doc, "数据库：", "熟悉 MySQL 表结构设计、事务、索引、锁、MVCC、联表查询与常见 SQL 优化方法。")
    add_skill_line(doc, "缓存/安全：", "熟悉 Redis 常用数据结构、缓存穿透/击穿/雪崩处理、分布式锁；掌握 JWT 认证机制。")
    add_skill_line(doc, "AI/工程化：", "了解 Spring AI、RAG、Function Calling、SSE 流式响应；熟悉 Maven、Git、Linux、Apifox、Docker 基础。")
    add_skill_line(doc, "前端联调：", "熟悉 HTML、CSS、JavaScript 与基础 Vue 开发，具备前后端联调和接口演示经验。")

    doc.core_properties.title = "何文旭-后端开发实习生简历"
    doc.core_properties.author = "何文旭"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
